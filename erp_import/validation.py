"""Lab 2 – Validation of the ERP product export (owner: Ami).

read_and_validate(path) checks an ERP export file BEFORE anything is saved:

* File level  -> the whole file is rejected (InvalidFileError):
  wrong file type, file can't be opened, required columns missing, no data rows.
* Row level   -> only that row is skipped, the other rows are still imported:
  missing Product ID / name, price not a number or negative,
  stock not a whole number or negative, Product ID used more than once.

It returns (rows, errors):
    rows   = [{"erp_id": "BOOK1001", "name": "1984", "description": "",
               "price": Decimal("9.99"), "stock": 25}, ...]   -> valid rows only
    errors = ["Row 4: Price 'abc' is not a number", ...]          -> one message per problem

The ERP (SAP Fiori "Export" button) creates an .xlsx file with the columns
Product ID | Product Name | Description | Price | Stock.
A .csv file with the same columns is accepted as well.
"""
import csv
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

ALLOWED_EXTENSIONS = {".xlsx", ".csv"}
MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB – an export of a few thousand products is far smaller
MAX_PRICE = Decimal("999999.99")  # Product.price has max_digits=8, decimal_places=2
MAX_ERP_ID_LENGTH = 50  # Product.erp_id max_length
MAX_NAME_LENGTH = 200  # Product.name max_length

# Column in the file -> field in the web shop.
# Several spellings are accepted (export labels and the ERP's technical names).
COLUMNS = {
    "erp_id": {"product id", "productid", "product_id", "id"},
    "name": {"product name", "productname", "name"},
    "description": {"description", "product description"},
    "price": {"price"},
    "stock": {"stock"},
}
REQUIRED_COLUMNS = ["erp_id", "name", "price", "stock"]  # description may be missing
COLUMN_LABELS = {"erp_id": "Product ID", "name": "Product Name", "description": "Description",
                 "price": "Price", "stock": "Stock"}


class InvalidFileError(ValueError):
    """The whole file can't be imported (wrong format, missing columns, empty...)."""


def read_and_validate(path):
    """Read an ERP export file and return (valid_rows, errors). See module docstring."""
    path = Path(path)
    _check_file(path)
    header, data_rows = _read_table(path)
    column_index = _map_columns(header)

    rows, errors = [], []
    seen_ids = {}  # erp_id -> first row number, to find duplicates
    for row_number, values in data_rows:
        if all(_is_empty(value) for value in values):
            continue  # ignore completely empty lines (e.g. at the end of a CSV)
        row, row_errors = _validate_row(row_number, values, column_index)
        if row_errors:
            errors.extend(row_errors)
            continue
        if row["erp_id"] in seen_ids:
            errors.append(f"Row {row_number}: Product ID '{row['erp_id']}' appears more than once "
                          f"(first in row {seen_ids[row['erp_id']]}) – row skipped")
            continue
        seen_ids[row["erp_id"]] = row_number
        rows.append(row)

    if not rows and not errors:
        raise InvalidFileError("The file contains no products.")
    return rows, errors


# ---------------------------------------------------------------- file level

def _check_file(path):
    if not path.is_file():
        raise InvalidFileError(f"File not found: {path.name}")
    if path.suffix.lower() not in ALLOWED_EXTENSIONS:
        raise InvalidFileError(f"Wrong file type '{path.suffix or 'none'}'. "
                               f"Please use the ERP export (.xlsx) or a .csv file.")
    size = path.stat().st_size
    if size == 0:
        raise InvalidFileError("The file is empty.")
    if size > MAX_FILE_SIZE:
        raise InvalidFileError("The file is too large (max. 5 MB).")


def _read_table(path):
    """Return (header, [(row_number, values), ...]). Row numbers match Excel/CSV line numbers."""
    if path.suffix.lower() == ".xlsx":
        return _read_xlsx(path)
    return _read_csv(path)


def _read_xlsx(path):
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, OSError, KeyError, ValueError) as exc:
        raise InvalidFileError("The file can't be opened as an Excel file. "
                               "Is it really an export from the ERP?") from exc
    try:
        sheet = workbook.worksheets[0]
        all_rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not all_rows:
        raise InvalidFileError("The file is empty.")
    header = list(all_rows[0])
    data = [(number, list(values)) for number, values in enumerate(all_rows[1:], start=2)]
    return header, data


def _read_csv(path):
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):  # UTF-8 (with/without BOM) or Windows Excel
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise InvalidFileError("The CSV file has an unknown text encoding.")
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    all_rows = list(csv.reader(text.splitlines(), dialect))
    if not all_rows:
        raise InvalidFileError("The file is empty.")
    data = [(number, values) for number, values in enumerate(all_rows[1:], start=2)]
    return all_rows[0], data


def _map_columns(header):
    """Find the position of each known column. Missing required columns -> InvalidFileError."""
    positions = {}
    for index, title in enumerate(header):
        title = str(title or "").strip().lower()
        for field, names in COLUMNS.items():
            if title in names and field not in positions:
                positions[field] = index
    missing = [COLUMN_LABELS[field] for field in REQUIRED_COLUMNS if field not in positions]
    if missing:
        raise InvalidFileError("Required column(s) missing: " + ", ".join(missing) +
                               ". Expected: Product ID, Product Name, Description, Price, Stock.")
    return positions


# ---------------------------------------------------------------- row level

def _validate_row(row_number, values, column_index):
    def cell(field):
        index = column_index.get(field)
        return values[index] if index is not None and index < len(values) else None

    errors = []
    prefix = f"Row {row_number}:"

    erp_id = _text(cell("erp_id"))
    if not erp_id:
        errors.append(f"{prefix} Product ID is missing")
    elif len(erp_id) > MAX_ERP_ID_LENGTH:
        errors.append(f"{prefix} Product ID is longer than {MAX_ERP_ID_LENGTH} characters")

    name = _text(cell("name"))
    if not name:
        errors.append(f"{prefix} Product Name is missing")
    elif len(name) > MAX_NAME_LENGTH:
        errors.append(f"{prefix} Product Name is longer than {MAX_NAME_LENGTH} characters")

    price, price_error = _parse_price(cell("price"))
    if price_error:
        errors.append(f"{prefix} {price_error}")

    stock, stock_error = _parse_stock(cell("stock"))
    if stock_error:
        errors.append(f"{prefix} {stock_error}")

    row = {"erp_id": erp_id, "name": name, "description": _text(cell("description")),
           "price": price, "stock": stock}
    return row, errors


def _parse_price(value):
    if _is_empty(value):
        return None, "Price is missing"
    if isinstance(value, bool):
        return None, f"Price '{value}' is not a number"
    if isinstance(value, (int, float, Decimal)):
        text = str(value)
    else:
        text = str(value).strip().replace("€", "").replace("EUR", "").strip()
        if "," in text and "." not in text:  # German decimal comma: 12,50
            text = text.replace(",", ".")
    try:
        price = Decimal(text)
    except InvalidOperation:
        return None, f"Price '{value}' is not a number"
    if not price.is_finite():
        return None, f"Price '{value}' is not a number"
    if price < 0:
        return None, f"Price {price} is negative"
    if price > MAX_PRICE:
        return None, f"Price {price} is too high (max. {MAX_PRICE})"
    return price.quantize(Decimal("0.01")), None


def _parse_stock(value):
    if _is_empty(value):
        return None, "Stock is missing"
    if isinstance(value, bool):
        return None, f"Stock '{value}' is not a whole number"
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        return None, f"Stock '{value}' is not a whole number"
    if not number.is_finite() or number != number.to_integral_value():
        return None, f"Stock '{value}' is not a whole number"
    if number < 0:
        return None, f"Stock {int(number)} is negative"
    return int(number), None


def _text(value):
    return "" if value is None else str(value).strip()


def _is_empty(value):
    return value is None or str(value).strip() == ""
